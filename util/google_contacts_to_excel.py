"""
Convierte exportaciones de Google Contactos (CSV) a un fichero Excel.
Mantiene los nombres de columna originales de Google Contactos.
Uso: python google_contacts_to_excel.py <fichero1.csv> [fichero2.csv ...] -o contactos.xlsx
"""

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
ALT_FILL    = PatternFill("solid", start_color="D6E4F0")
NORMAL_FILL = PatternFill("solid", start_color="FFFFFF")
DATA_FONT   = Font(name="Arial", size=10)
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT        = Alignment(horizontal="left",   vertical="center", wrap_text=True)


def load_csv(path: Path, source_name: str) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str).fillna("")
    df.insert(0, "Fuente", source_name)
    # Eliminar filas sin datos útiles
    name_cols  = [c for c in df.columns if "Name" in c and "Phonetic" not in c and "Label" not in c]
    phone_cols = [c for c in df.columns if "Phone" in c and "Value" in c]
    email_cols = [c for c in df.columns if "E-mail" in c and "Value" in c]
    check_cols = name_cols + phone_cols + email_cols
    mask = df[check_cols].apply(lambda row: row.str.strip().ne("").any(), axis=1)
    df = df[mask].reset_index(drop=True)
    # Eliminar columnas completamente vacías (excepto "Fuente")
    non_empty = [c for c in df.columns if c == "Fuente" or df[c].str.strip().ne("").any()]
    return df[non_empty]


def write_sheet(ws, df: pd.DataFrame):
    cols = list(df.columns)

    # Encabezado
    ws.append(cols)
    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = CENTER

    # Datos
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        fill = ALT_FILL if row_idx % 2 == 0 else NORMAL_FILL
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")
            cell.fill      = fill
            cell.font      = DATA_FONT
            cell.alignment = LEFT

    # Ancho de columnas adaptado al nombre del encabezado
    for col_idx, col_name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(col_name) + 2, 14)

    ws.row_dimensions[1].height = 22
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    ws.freeze_panes    = "B2"


def build_excel(frames: dict, output_path: Path):
    wb = Workbook()
    wb.remove(wb.active)

    # Hoja combinada: unión de todas las columnas
    all_df = pd.concat(list(frames.values()), ignore_index=True).fillna("")
    ws_all = wb.create_sheet("Todos")
    write_sheet(ws_all, all_df)

    # Una hoja por fichero con sus columnas originales
    for name, df in frames.items():
        ws = wb.create_sheet(name[:31])
        write_sheet(ws, df)

    wb.save(output_path)
    total = sum(len(d) for d in frames.values())
    print(f"✅  Excel guardado en: {output_path}  ({total} contactos)")


def main():
    parser = argparse.ArgumentParser(description="Google Contacts CSV → Excel")
    parser.add_argument("csvs", nargs="+", help="Ficheros CSV exportados de Google Contactos")
    parser.add_argument("-o", "--output", default="contactos.xlsx", help="Fichero Excel de salida")
    args = parser.parse_args()

    frames = {}
    for csv_path_str in args.csvs:
        p = Path(csv_path_str)
        if not p.exists():
            print(f"⚠️  No encontrado: {p}", file=sys.stderr)
            continue
        name = p.stem.replace("contacts_", "").replace("_", " ").title()
        df   = load_csv(p, name)
        frames[name] = df
        print(f"   {p.name}: {len(df)} contactos  ({len(df.columns)-1} columnas originales)")

    if not frames:
        print("❌  No se cargaron datos.", file=sys.stderr)
        sys.exit(1)

    build_excel(frames, Path(args.output))


if __name__ == "__main__":
    main()
